"""이전 사이트 실제 회원/구매평 데이터 이관.

- 기존 더미 유저(admin 제외) 전부 삭제 (CASCADE 로 연관 주문/수강/문서도 정리됨)
- member_data.xlsx (2675 행) 를 실제 회원으로 import
  - username = 고유키 (영숫자/언더스코어만이라 그대로 사용 가능, 검증 완료)
  - password_hash = None → 로그인은 "비밀번호 재설정" 이메일 플로우로 유도
  - is_verified = False (이메일 실제 수신 여부 재검증 필요)
  - created_at 은 원래 가입일 그대로 유지 (KST 로 해석)
- 기존 더미 후기(Post, category=review) 전부 삭제
- reviews.xlsx (158 행) 를 실제 후기로 import — 상품명은 old-site 표기라
  PRODUCT_TITLE_MAP 으로 현재 course.title 에 매핑

멱등하지 않음 — 한 번만 실행할 것.

실행: `cd backend && venv/bin/python scripts/import_legacy_data.py`
"""

from __future__ import annotations

import re
import sys
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl  # noqa: E402
from sqlalchemy import select  # noqa: E402

from app.core.database import SessionLocal  # noqa: E402
from app.models.community import Post, PostCategory  # noqa: E402
from app.models.course import Course  # noqa: E402
from app.models.user import User  # noqa: E402

KST = timezone(timedelta(hours=9))

DATA_DIR = Path(
    "/private/tmp/claude-501/-Users-kimnahyun-github-kcpec-platform/"
    "3c17a1ea-d8e8-4cbe-9392-4d8cab2f7dd2/scratchpad/customer_import"
)

# 이전 사이트 상품명(구매평 데이터 기준) → 현재 course.title.
# "올인원" 번들은 상담이 아니라 강의 자체에 대한 후기로 보고 기반 강의에 매핑.
PRODUCT_TITLE_MAP: dict[str, str] = {
    "준법의식 강화": "준법의식 강화",
    "디지털 성범죄 예방": "디지털 성범죄 예방",
    "성범죄 재범 방지강의": "성범죄 예방",
    "음주운전 예방 강의": "음주운전 예방",
    "스토킹": "스토킹범죄 예방",
    "성매매": "성매매 예방",
    "학교폭력": "학교폭력 예방",
    "음주운전 재범예방 올인원(음주운전강의+준법의식강의+전문가상담)": "음주운전 예방",
    "성범죄 재범예방 올인원(성범죄강의+준법의식강의+전문가상담)": "성범죄 예방",
    "도박 및 도박개장": "도박 및 도박개장 예방",
    "마약": "마약 예방",
    "준법의식 올인원(준법의식강의+전문가상담)": "준법의식 강화",
    "성매매 재범예방 올인원(성매매강의+준법의식강의+전문가상담)": "성매매 예방",
    "디지털 성범죄 예방 올인원(디지털 성범죄 예방강의+준법의식강의+전문가상담)": "디지털 성범죄 예방",
    "사기횡령배임 등 재산범죄": "사기횡령배임 등 재산범죄 예방",
    "전문가 심리상담 프로그램": "기본 프로그램",
    "도박 및 도박개장 예방 올인원(도박 및 도박개장 예방강의+준법의식강의+전문가상담)": "도박 및 도박개장 예방",
    "스토킹범죄 올인원(스토킹범죄예방강의+준법의식강의+전문가상담)": "스토킹범죄 예방",
}


def parse_kst(s: str | None) -> datetime | None:
    if not s:
        return None
    s = str(s).strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt).replace(tzinfo=KST)
        except ValueError:
            continue
    return None


def parse_date(s: str | None) -> date | None:
    if not s:
        return None
    try:
        return datetime.strptime(str(s).strip(), "%Y-%m-%d").date()
    except ValueError:
        return None


def import_members(db) -> int:
    wb = openpyxl.load_workbook(DATA_DIR / "member_data.xlsx", data_only=True)
    ws = wb["Sheet1"]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {h: i for i, h in enumerate(headers)}

    count = 0
    seen_usernames: set[str] = set()
    seen_emails: set[str] = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        key = str(row[idx["고유키"]]).strip()
        email = str(row[idx["이메일"]]).strip()
        username = re.sub(r"[^A-Za-z0-9_]", "_", key)[:50]
        if username in seen_usernames or email in seen_emails:
            continue  # 방어적 스킵 (사전 스캔에서는 없었지만 혹시 몰라)
        seen_usernames.add(username)
        seen_emails.add(email)

        name = str(row[idx["이름"]]).strip() or None
        phone_raw = row[idx["연락처"]]
        phone = str(phone_raw).strip() or None if phone_raw else None
        birth = parse_date(row[idx["생년월일"]])
        created = parse_kst(row[idx["가입일"]]) or datetime.now(KST)

        user = User(
            username=username,
            password_hash=None,
            email=email,
            name=name,
            phone=phone,
            birth_date=birth,
            is_active=True,
            is_admin=False,
            is_verified=False,
            created_at=created,
        )
        db.add(user)
        count += 1
    return count


def import_reviews(db) -> tuple[int, int]:
    course_by_title = {c.title: c for c in db.scalars(select(Course)).all()}

    wb = openpyxl.load_workbook(DATA_DIR / "reviews.xlsx", data_only=True)
    ws = wb["Sheet1"]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {h: i for i, h in enumerate(headers)}

    count = 0
    skipped = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        old_title = row[idx["상품명"]]
        new_title = PRODUCT_TITLE_MAP.get(old_title)
        course = course_by_title.get(new_title) if new_title else None
        if course is None:
            skipped += 1
            continue

        content = str(row[idx["글 내용"]]).strip()
        author = str(row[idx["작성자"]]).strip() or "익명"
        rating = int(row[idx["평점"]] or 5)
        created = parse_kst(row[idx["작성시각"]]) or datetime.now(KST)

        post = Post(
            title=content[:255],
            content=content,
            category=PostCategory.REVIEW,
            author_name=author,
            course_id=course.id,
            course_category=course.category.value,
            rating=max(1, min(5, rating)),
            created_at=created,
        )
        db.add(post)
        count += 1
    return count, skipped


def main() -> None:
    db = SessionLocal()
    try:
        deleted_users = db.query(User).filter(User.is_admin.is_(False)).delete(
            synchronize_session=False
        )
        deleted_reviews = (
            db.query(Post)
            .filter(Post.category == PostCategory.REVIEW)
            .delete(synchronize_session=False)
        )
        db.commit()
        print(f"삭제: 더미 유저 {deleted_users}명, 더미 후기 {deleted_reviews}건")

        member_count = import_members(db)
        db.commit()
        print(f"회원 import: {member_count}명")

        review_count, review_skipped = import_reviews(db)
        db.commit()
        print(f"후기 import: {review_count}건 (매핑 실패로 스킵: {review_skipped}건)")
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


if __name__ == "__main__":
    main()
