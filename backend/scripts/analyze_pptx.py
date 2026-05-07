"""PPTX 파일의 모든 슬라이드 / shape 의 텍스트를 출력하는 분석 스크립트.

사용:
    python scripts/analyze_pptx.py "static/templates/certificates/1. 준법의식 수료증.pptx"
    python scripts/analyze_pptx.py --all   # certificates 폴더의 모든 pptx 일괄 분석
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation


def _walk_shape(shape, indent: int = 0) -> None:
    pad = "  " * indent
    base = (
        f"{pad}Shape '{shape.name}' (type={shape.shape_type}, "
        f"top={shape.top}, left={shape.left}, "
        f"width={shape.width}, height={shape.height})"
    )

    # Group shape — recurse into children
    if shape.shape_type == 6:  # GROUP
        print(base + " [GROUP]")
        for child in shape.shapes:
            _walk_shape(child, indent + 1)
        return

    # Table — iterate cells
    if shape.has_table:
        print(base + " [TABLE]")
        table = shape.table
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                txt = cell.text_frame.text if cell.text_frame else ""
                print(f"{pad}  [{r_idx},{c_idx}] {repr(txt[:200])}")
        return

    # Plain text frame
    if shape.has_text_frame:
        print(f"{base}: {repr(shape.text_frame.text[:200])}")
        return

    # Picture / etc — just position info
    print(base)


def analyze(pptx_path: Path) -> None:
    print(f"\n##### {pptx_path.name} #####")
    prs = Presentation(str(pptx_path))
    for slide_num, slide in enumerate(prs.slides, start=1):
        print(f"\n=== Slide {slide_num} ===")
        for shape in slide.shapes:
            _walk_shape(shape)


def main() -> None:
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)

    arg = sys.argv[1]
    cert_dir = Path(__file__).resolve().parents[1] / "static" / "templates" / "certificates"

    if arg == "--all":
        targets = sorted(cert_dir.glob("*.pptx"))
    else:
        p = Path(arg)
        targets = [p if p.is_absolute() else cert_dir / p.name]

    for t in targets:
        if not t.exists():
            print(f"[skip] {t} not found")
            continue
        analyze(t)


if __name__ == "__main__":
    main()
