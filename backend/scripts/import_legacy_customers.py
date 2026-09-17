"""구 사이트 실 회원 + 주문/결제 이력 이관 (2차, 2026-09).

`import_legacy_data.py`(9/6, 회원+후기 이관용)와 달리 이 스크립트는
**기존 행을 절대 삭제하지 않고 오직 추가(INSERT)만 한다** — 지금 새
플랫폼에 실사용자가 이미 실주문/수강/발급서류를 갖고 있어서, 기존
스크립트처럼 "admin 제외 전체 유저 삭제" 하는 방식은 여기선 쓸 수 없다.

입력 파일 (전부 고객 PII 포함 — 서버에서 실행 후 반드시 삭제할 것):
- MEMBER_XLSX: 회원 명부 (고유키/이메일/이름/연락처/생년월일/가입일 등)
- ORDER_XLSX : 주문내역 (주문번호/주문자 이메일/상품명/품목실결제가/주문일 등)
- PAYMENT_XLSX: 결제내역 (주문번호/결제수단/결제상태/결제시간/PG거래번호 등,
  ORDER_XLSX 보다 오래된 날짜라 최근 주문 일부는 매칭 안 될 수 있음 — 그 경우
  결제수단은 최빈값(카드)으로, 상태는 주문 자체의 "거래종료" 여부로 추정)

처리 순서: 회원 이관(이메일 중복이면 기존 계정 재사용) → 주문번호 단위로
묶어서 Order 생성(품목별 1행, 2개 이상이면 bundle_id 로 묶음) → 결제완료
COURSE 주문 중 **주문일이 스크립트 실행 시점 기준 7일 이내인 것만**
Enrollment 생성(enrolled_at=지금, expires_at=지금+7일, is_completed=False
— "구매 이력만" 가져오고 수강 완료 처리는 안 함). 7일보다 오래된 주문은
Order(구매 이력)만 남기고 Enrollment 는 아예 만들지 않음 — 마이페이지
"내 강의" 에는 안 뜨고 "주문내역"에만 보임. (2026-09-17 확정, 새 구매자에게
적용되는 1개월(30일) 규정과는 별개로 이번 이관 건에서만 예외적으로 7일
기준 사용.)

멱등하지 않음 — 이미 생성된 이메일은 재실행 시 건너뛰지만(회원 단계),
주문/Enrollment 단계는 재실행 시 중복 생성될 수 있으므로 한 번만 실행할 것.

실행:
  cd backend && venv/bin/python scripts/import_legacy_customers.py          # dry-run(커밋 안 함, 요약만 출력)
  cd backend && venv/bin/python scripts/import_legacy_customers.py --commit # 실제 반영
"""

from __future__ import annotations

import re
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl  # noqa: E402
from sqlalchemy import select  # noqa: E402

from app.core.database import SessionLocal  # noqa: E402
from app.models.course import Course, CourseCategory  # noqa: E402
from app.models.enrollment import Enrollment  # noqa: E402
from app.models.order import Order, OrderStatus, OrderType, PaymentMethod  # noqa: E402
from app.models.user import User  # noqa: E402

KST = timezone(timedelta(hours=9))

DATA_DIR = Path(
    "/private/tmp/claude-501/-Users-kimnahyun-github-kcpec-platform/"
    "3c17a1ea-d8e8-4cbe-9392-4d8cab2f7dd2/scratchpad/member_import"
)
MEMBER_XLSX = DATA_DIR / "member.xlsx"
ORDER_XLSX = DATA_DIR / "orders.xlsx"
PAYMENT_XLSX = DATA_DIR / "payments.xlsx"

# 구 사이트 상품명(주문내역 기준) → 현재 course.title.
# "올인원" 묶음 상품은 번들 구성요소별로 쪼개지 않고 주 강의 하나로만 매핑한다
# (범위를 "고객이 뭘 샀었는지 보여주기"로 좁혀 리스크를 줄임).
PRODUCT_TITLE_MAP: dict[str, str] = {
    "준법의식 강화": "준법의식 강화",
    "디지털 성범죄 예방": "디지털 성범죄 예방",
    "성범죄 재범 방지강의": "성범죄 예방",
    "음주운전 예방 강의": "음주운전 예방",
    "스토킹": "스토킹범죄 예방",
    "성매매": "성매매 예방",
    "학교폭력": "학교폭력 예방",
    "피싱범죄": "피싱범죄 예방",
    "마약": "마약 예방",
    "도박 및 도박개장": "도박 및 도박개장 예방",
    "사기횡령배임 등 재산범죄": "사기횡령배임 등 재산범죄 예방",
    "전문가 심리상담 프로그램": "기본 프로그램",
    "음주운전 재범예방 올인원(음주운전강의+준법의식강의+전문가상담)": "음주운전 예방",
    "성범죄 재범예방 올인원(성범죄강의+준법의식강의+전문가상담)": "성범죄 예방",
    "준법의식 올인원(준법의식강의+전문가상담)": "준법의식 강화",
    "성매매 재범예방 올인원(성매매강의+준법의식강의+전문가상담)": "성매매 예방",
    "디지털 성범죄 예방 올인원(디지털 성범죄 예방강의+준법의식강의+전문가상담)": "디지털 성범죄 예방",
    "도박 및 도박개장 예방 올인원(도박 및 도박개장 예방강의+준법의식강의+전문가상담)": "도박 및 도박개장 예방",
    "스토킹범죄 올인원(스토킹범죄예방강의+준법의식강의+전문가상담)": "스토킹범죄 예방",
    "마약 재범예방 올인원(마약강의+준법의식강의+전문가상담)": "마약 예방",
    "사기횡령배임 등 재산범죄 올인원(재산범죄예방강의+준법의식강의+전문가상담)": "사기횡령배임 등 재산범죄 예방",
    "피싱범죄예방 올인원(피싱범죄예방강의+준법의식강의+전문가상담)": "피싱범죄 예방",
}

PAYMENT_METHOD_MAP: dict[str, PaymentMethod] = {
    "카드": PaymentMethod.CARD,
    "실시간계좌이체": PaymentMethod.TRANSFER,
    "무통장입금": PaymentMethod.BANK_TRANSFER,
    "삼성페이": PaymentMethod.SAMSUNGPAY,
    "가상계좌": PaymentMethod.BANK_TRANSFER,
}

PAYMENT_STATUS_MAP: dict[str, OrderStatus] = {
    "결제완료": OrderStatus.PAID,
    "부분환불": OrderStatus.PAID,
    "전체환불": OrderStatus.REFUNDED,
    "입금전 취소": OrderStatus.CANCELLED,
    "결제기한초과": OrderStatus.CANCELLED,
}

# 이관 시점 기준 이 기간 이내에 구매한 경우에만 강의 접근을 열어준다(Enrollment
# 생성). 그보다 오래된 주문은 Order(구매 이력)만 남기고 접근은 안 열어준다.
# 신규 결제자에게 적용되는 사이트 전역 설정(현재 30일, ENROLLMENT_ACCESS_DAYS)
# 과는 별개로, 이번 레거시 이관 건에 한해 예전 정책(7일)을 그대로 씀
# (2026-09-17 확정).
LEGACY_RECENT_CUTOFF_DAYS = 7
LEGACY_ACCESS_WINDOW_DAYS = 7


def parse_kst(s: str | None) -> datetime | None:
    if not s:
        return None
    s = str(s).strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt).replace(tzinfo=KST)
        except ValueError:
            continue
    return None


def parse_date(s: str | None) -> date | None:
    if not s:
        return None
    s = str(s).strip()
    if s in ("0000-00-00", ""):
        return None
    try:
        return datetime.strptime(s, "%Y-%m-%d").date()
    except ValueError:
        return None


def sheet_rows(path: Path, sheet: str | None = None) -> tuple[dict[str, int], list[tuple]]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {h: i for i, h in enumerate(headers)}
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    return idx, rows


def import_members(db, commit: bool) -> tuple[dict[str, int], int, int]:
    """이메일 → user_id 매핑을 반환. 이미 있는 이메일은 기존 user_id 재사용."""
    idx, rows = sheet_rows(MEMBER_XLSX)

    existing = {
        (u.email or "").strip().lower(): u.id
        for u in db.scalars(select(User)).all()
    }
    existing_usernames = {u.username for u in db.scalars(select(User)).all()}

    email_to_user_id: dict[str, int] = dict(existing)
    created = 0
    reused = 0

    for row in rows:
        email = str(row[idx["이메일"]] or "").strip().lower()
        if not email:
            continue
        if email in existing:
            reused += 1
            continue  # 이미 실사용자로 존재 — 새로 안 만들고 그 user_id 그대로 재사용

        key = str(row[idx["고유키"]]).strip()
        username = re.sub(r"[^A-Za-z0-9_]", "_", key)[:50]
        base_username = username
        suffix = 1
        while username in existing_usernames:
            suffix += 1
            username = f"{base_username[:46]}_{suffix}"
        existing_usernames.add(username)

        name = str(row[idx["이름"]] or "").strip() or None
        phone_raw = row[idx["연락처"]]
        phone = str(phone_raw).strip() or None if phone_raw else None
        birth = parse_date(row[idx["생년월일"]])
        created_at = parse_kst(row[idx["가입일"]]) or datetime.now(KST)

        user = User(
            username=username,
            password_hash=None,
            email=email,
            name=name,
            phone=phone,
            birth_date=birth,
            is_active=True,
            is_admin=False,
            is_verified=False,
            created_at=created_at,
        )
        db.add(user)
        db.flush()  # id 확보
        email_to_user_id[email] = user.id
        created += 1

    if commit:
        db.commit()
    return email_to_user_id, created, reused


def import_orders_and_enrollments(
    db, email_to_user_id: dict[str, int], commit: bool
) -> dict[str, int]:
    order_idx, order_rows = sheet_rows(ORDER_XLSX)
    pay_idx, pay_rows = sheet_rows(PAYMENT_XLSX)

    # 주문번호 → 결제내역 매칭 (결제완료가 있으면 그것 우선, 없으면 최신 1건)
    pay_by_order: dict[str, dict] = {}
    for r in pay_rows:
        onum = str(r[pay_idx["주문번호"]])
        info = {
            "결제수단": r[pay_idx["결제수단"]],
            "결제상태": r[pay_idx["결제상태"]],
            "결제시간": r[pay_idx["결제시간"]],
            "PG거래번호": r[pay_idx["PG거래번호"]],
        }
        prev = pay_by_order.get(onum)
        if prev is None or info["결제상태"] == "결제완료":
            pay_by_order[onum] = info

    courses_by_title = {c.title: c for c in db.scalars(select(Course)).all()}

    # 주문번호 단위로 품목 묶기
    groups: dict[str, list[tuple]] = defaultdict(list)
    for r in order_rows:
        groups[str(r[order_idx["주문번호"]])].append(r)

    now = datetime.now(KST)
    recent_cutoff = now - timedelta(days=LEGACY_RECENT_CUTOFF_DAYS)

    stats = {
        "orders_created": 0,
        "orders_skipped_no_user": 0,
        "orders_skipped_no_course": 0,
        "orders_skipped_no_amount": 0,
        "payment_defaulted": 0,
        "enrollments_created": 0,
        "enrollments_skipped_dup": 0,
        "enrollments_skipped_too_old": 0,
        "status_paid": 0,
        "status_refunded": 0,
        "status_cancelled": 0,
    }

    for onum, items in groups.items():
        buyer_email = str(items[0][order_idx["주문자 이메일"]] or "").strip().lower()
        user_id = email_to_user_id.get(buyer_email)
        if user_id is None:
            stats["orders_skipped_no_user"] += len(items)
            continue

        bundle_id = f"legacy-{onum}" if len(items) > 1 else None
        pay = pay_by_order.get(onum)
        if pay is not None:
            payment_method = PAYMENT_METHOD_MAP.get(pay["결제수단"], PaymentMethod.CARD)
            status = PAYMENT_STATUS_MAP.get(pay["결제상태"], OrderStatus.PAID)
            paid_at = parse_kst(pay["결제시간"]) if status == OrderStatus.PAID else None
            toss_payment_key = pay["PG거래번호"] or None
        else:
            stats["payment_defaulted"] += 1
            payment_method = PaymentMethod.CARD
            order_status_raw = items[0][order_idx["주문상태"]]
            status = OrderStatus.PAID if order_status_raw == "거래종료" else OrderStatus.CANCELLED
            created_dt = parse_kst(items[0][order_idx["주문일"]])
            paid_at = created_dt if status == OrderStatus.PAID else None
            toss_payment_key = None

        for item in items:
            product_name = item[order_idx["상품명"]]
            course_title = PRODUCT_TITLE_MAP.get(product_name)
            course = courses_by_title.get(course_title) if course_title else None
            if course is None:
                stats["orders_skipped_no_course"] += 1
                continue

            amount = item[order_idx["품목실결제가"]]
            if amount is None:
                stats["orders_skipped_no_amount"] += 1
                continue

            created_at = parse_kst(item[order_idx["주문일"]]) or datetime.now(KST)
            order_type = (
                OrderType.COUNSELING if course.category == CourseCategory.COUNSELING else OrderType.COURSE
            )

            order = Order(
                user_id=user_id,
                course_id=course.id,
                order_type=order_type,
                toss_payment_key=toss_payment_key,
                payment_method=payment_method,
                amount=int(amount),
                status=status,
                bundle_id=bundle_id,
                created_at=created_at,
                paid_at=paid_at,
            )
            db.add(order)
            stats["orders_created"] += 1
            if status == OrderStatus.PAID:
                stats["status_paid"] += 1
            elif status == OrderStatus.REFUNDED:
                stats["status_refunded"] += 1
            else:
                stats["status_cancelled"] += 1

            order_date = paid_at or created_at
            if order_type == OrderType.COURSE and status == OrderStatus.PAID:
                if order_date < recent_cutoff:
                    # 이관 시점 기준 7일보다 오래된 구매 — 주문 이력만 남기고
                    # 강의 접근(Enrollment)은 열어주지 않는다.
                    stats["enrollments_skipped_too_old"] += 1
                    continue
                dup = db.scalar(
                    select(Enrollment).where(
                        Enrollment.user_id == user_id, Enrollment.course_id == course.id
                    )
                )
                if dup is not None:
                    stats["enrollments_skipped_dup"] += 1
                else:
                    db.add(
                        Enrollment(
                            user_id=user_id,
                            course_id=course.id,
                            # 지금 시점부터 새로 7일간 접근 가능하도록 오픈
                            # (원래 예전 구매일 기준이 아니라 "지금 다시 열어줌").
                            enrolled_at=now,
                            expires_at=now + timedelta(days=LEGACY_ACCESS_WINDOW_DAYS),
                            is_completed=False,
                        )
                    )
                    stats["enrollments_created"] += 1

    if commit:
        db.commit()
    return stats


def main() -> None:
    commit = "--commit" in sys.argv
    db = SessionLocal()
    try:
        email_to_user_id, created, reused = import_members(db, commit)
        print(f"회원: 신규 생성 {created}명 / 기존 계정 재사용 {reused}명")

        stats = import_orders_and_enrollments(db, email_to_user_id, commit)
        for k, v in stats.items():
            print(f"  {k}: {v}")

        if commit:
            print("\n[커밋 완료]")
        else:
            db.rollback()
            print("\n[DRY RUN — 실제 반영 안 됨. --commit 옵션으로 재실행하면 반영됩니다]")
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


if __name__ == "__main__":
    main()
