"""심리상담 의견서 DOCX 생성기.

새 템플릿(9 rows × 6 cols) 기준 셀 매핑:
  R0 C0 : 증번호 (cols 0-2 merged)
  R3 C1 : 성명
  R3 C3 : 성별
  R3 C5 : 생년월일
  R4 C1 : 상담일자 (cols 1-3 merged)
  R4 C5 : 연락처
  R5 C1 : 상담배경 및 상담취지 (cols 1-5 merged)
  R6 C1 : 상담내용 및 종합소견 (cols 1-5 merged)
  R8 C0 : 발급일 (cell 의 "년/월/일" paragraph 만 갱신)

설문 응답(survey.responses)은 두 가지 형태를 모두 받아들인다:
  - 신형: { "personal": {name, gender, birthdate, phone, ...}, "q2"..."q6": "..." }
  - 구형: { "인적사항": "텍스트", "사건내용": "...", ... }  (regex 폴백)
"""

from __future__ import annotations

import re
import shutil
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from docx import Document

from app.core.docx_utils import set_cell_text, set_paragraph_text

from app.models.counseling import CounselingSurvey
from app.models.user import User

BACKEND_DIR = Path(__file__).resolve().parents[2]
TEMPLATE_PATH = (
    BACKEND_DIR / "static" / "templates" / "counseling_template.docx"
)


# ---------- 셀 텍스트 치환 ----------------------------------------------------


# 표/문단 텍스트 치환 헬퍼는 app.core.docx_utils 로 옮겼다(반성문·탄원서
# 생성기와 공용). 기존 호출부 호환을 위해 이름만 로컬로 다시 노출.
_set_cell_text = set_cell_text
_set_paragraph_text = set_paragraph_text


# ---------- 인적사항 추출 -----------------------------------------------------


_AGE_RE = re.compile(r"(\d{1,3})\s*세")
_PHONE_RE = re.compile(r"(\d{2,3}[-\s]?\d{3,4}[-\s]?\d{4})")
_BIRTH_RE = re.compile(
    r"(\d{2,4})\s*[년./-]\s*(\d{1,2})\s*[월./-]\s*(\d{1,2})\s*일?"
)


@dataclass
class PersonalSnapshot:
    name: str = ""
    gender: str = ""
    birthdate: str = ""   # 표시용: "1991년 3월 15일" 또는 입력 문자열 그대로
    phone: str = ""


def _format_korean_birth(s: str) -> str:
    """ISO 'YYYY-MM-DD' 또는 'YY년 MM월 DD일' 등 다양한 형식을
    "YYYY년 M월 D일" 한글 형식으로 정규화. 못 알아보면 원문 반환.
    """
    if not s:
        return ""
    s = s.strip()
    # ISO YYYY-MM-DD
    iso_m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", s)
    if iso_m:
        y, mo, d = iso_m.groups()
        return f"{int(y)}년 {int(mo)}월 {int(d)}일"
    bm = _BIRTH_RE.match(s)
    if bm:
        y, mo, d = bm.groups()
        if len(y) == 2:
            y = f"19{y}" if int(y) > 30 else f"20{y}"
        return f"{int(y)}년 {int(mo)}월 {int(d)}일"
    return s


def _parse_legacy_personal(q1_text: str, user: User) -> PersonalSnapshot:
    """구형 데이터: q1(인적사항) 자유 텍스트에서 정규식으로 보완 추출."""
    name = user.username or ""
    gender = ""
    birthdate = ""
    phone = ""
    if "남" in q1_text and "남자" not in q1_text and "여자" not in q1_text:
        gender = "남"
    elif "남자" in q1_text or "남성" in q1_text:
        gender = "남"
    elif "여자" in q1_text or "여성" in q1_text:
        gender = "여"
    if user.birth_date:
        birthdate = (
            f"{user.birth_date.year}년 {user.birth_date.month}월 "
            f"{user.birth_date.day}일"
        )
    else:
        bm = _BIRTH_RE.search(q1_text)
        if bm:
            y, mo, d = bm.groups()
            if len(y) == 2:
                y = f"19{y}" if int(y) > 30 else f"20{y}"
            birthdate = f"{int(y)}년 {int(mo)}월 {int(d)}일"
    pm = _PHONE_RE.search(q1_text)
    if pm:
        phone = pm.group(1).replace(" ", "-")
    return PersonalSnapshot(name=name, gender=gender, birthdate=birthdate, phone=phone)


def _personal_snapshot(survey: CounselingSurvey, user: User) -> PersonalSnapshot:
    """survey.responses 에서 personal dict 추출 — 신형 우선, 구형은 폴백."""
    responses: dict[str, Any] = survey.responses or {}
    p = responses.get("personal")
    if isinstance(p, dict):
        return PersonalSnapshot(
            name=str(p.get("name", "") or "").strip() or (user.username or ""),
            gender=str(p.get("gender", "") or "").strip(),
            birthdate=_format_korean_birth(str(p.get("birthdate", "") or "")),
            phone=str(p.get("phone", "") or "").strip(),
        )
    legacy_q1 = responses.get("인적사항") or responses.get("q1") or ""
    return _parse_legacy_personal(str(legacy_q1), user)


# ---------- draft_text 섹션 분리 ---------------------------------------------


def _split_draft_sections(draft_text: str) -> tuple[str, str]:
    if "[상담배경]" in draft_text and "[상담내용]" in draft_text:
        bg = (
            draft_text.split("[상담배경]", 1)[1]
            .split("[상담내용]", 1)[0]
            .strip()
        )
        ct = draft_text.split("[상담내용]", 1)[1].strip()
        return bg, ct
    return "", draft_text.strip()


# ---------- 메인 생성 함수 ----------------------------------------------------


def build_doc_number(survey_id: int, today: date) -> str:
    return f"증 {today.year}-kcpec-{today.month:02d}-{survey_id:05d} 호"


def fill_counseling_template(
    survey: CounselingSurvey,
    user: User,
    draft_text: str,
    output_path: Path,
    today: date | None = None,
) -> Path:
    """양식에 데이터 채워서 output_path 에 DOCX 저장."""
    if not TEMPLATE_PATH.exists():
        raise RuntimeError(f"템플릿 파일 누락: {TEMPLATE_PATH}")

    today = today or datetime.now().date()
    info = _personal_snapshot(survey, user)
    background, content = _split_draft_sections(draft_text)
    doc_number = build_doc_number(survey.id, today)
    counseling_date = f"{today.year}년 {today.month}월 {today.day}일"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(TEMPLATE_PATH, output_path)

    doc = Document(str(output_path))
    table = doc.tables[0]

    _set_cell_text(table.rows[0].cells[0], doc_number)
    _set_cell_text(table.rows[3].cells[1], info.name)
    _set_cell_text(table.rows[3].cells[3], info.gender)
    _set_cell_text(table.rows[3].cells[5], info.birthdate)
    _set_cell_text(table.rows[4].cells[1], counseling_date)
    _set_cell_text(table.rows[4].cells[5], info.phone)
    _set_cell_text(
        table.rows[5].cells[1], background or "(설문 응답 기반 초안 없음)"
    )
    _set_cell_text(table.rows[6].cells[1], content or "(초안 없음)")

    # 발급일: R8 셀의 paragraph 들 중 "년/월/일" 패턴을 가진 것을 찾아 갱신.
    # (전체 셀을 비우면 가운데 정렬·여백 같은 템플릿 레이아웃이 깨짐)
    issued_str = f"{today.year}년   {today.month}월   {today.day}일"
    issue_cell = table.rows[8].cells[0]
    for para in issue_cell.paragraphs:
        if "년" in para.text and "월" in para.text and "일" in para.text:
            _set_paragraph_text(para, issued_str)
            break

    doc.save(str(output_path))
    return output_path


__all__ = [
    "build_doc_number",
    "fill_counseling_template",
]
