"""DOCX 표/문단 텍스트 치환 공용 헬퍼.

document_generator.py(심리상담 의견서)와 legal_letter_generator.py(반성문·
탄원서)가 같은 방식으로 템플릿의 표 셀/문단 텍스트를 갈아치우므로 공용
모듈로 분리했다 — 서식/폰트는 보존하고 텍스트만 바꾼다.
"""

from __future__ import annotations

from docx.table import _Cell


def set_cell_text(cell: _Cell, text: str) -> None:
    """셀 내용을 text 로 교체. 첫 paragraph 첫 run 의 폰트/크기를 보존하고
    멀티라인은 paragraph 분리.
    """
    paragraphs = list(cell.paragraphs)
    lines = (text or "").split("\n") if text else [""]

    first_p = paragraphs[0] if paragraphs else cell.add_paragraph()
    template_run = first_p.runs[0] if first_p.runs else None
    template_font_name = template_run.font.name if template_run else None
    template_font_size = template_run.font.size if template_run else None

    # 첫 run 의 텍스트만 첫 라인으로 갈아치우고 나머지 run 텍스트는 비움.
    if template_run is not None:
        template_run.text = lines[0]
        for r in first_p.runs[1:]:
            r.text = ""
    else:
        first_p.add_run(lines[0])

    # 첫 paragraph 외 paragraph 제거.
    for p in paragraphs[1:]:
        p._element.getparent().remove(p._element)

    # 추가 라인은 새 paragraph 로 (스타일/폰트 모방).
    for line in lines[1:]:
        p = cell.add_paragraph()
        run = p.add_run(line)
        if template_font_name:
            run.font.name = template_font_name
        if template_font_size:
            run.font.size = template_font_size


def set_paragraph_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for r in paragraph.runs[1:]:
            r.text = ""
    else:
        paragraph.add_run(text)
